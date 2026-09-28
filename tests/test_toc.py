import docx
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from handbook_generator import toc
from handbook_generator.toc import locate, toc_entries, update_toc_numbers


def _entry_xml(number_prefix, title, number_runs):
    runs = "".join(f"<w:r><w:t xml:space=\"preserve\">{n}</w:t></w:r>" for n in number_runs)
    return (
        f"<w:p><w:r><w:t>{number_prefix}</w:t></w:r><w:r><w:tab/></w:r>"
        f"<w:r><w:t xml:space=\"preserve\">{title}</w:t></w:r><w:r><w:tab/></w:r>{runs}</w:p>"
    )


def _doc_with_toc(tmp_path, entries):
    d = docx.Document()
    d.add_paragraph("Cover")
    paras = "".join(_entry_xml(*e) for e in entries)
    d.element.body.insert(1, parse_xml(f"<w:sdt {nsdecls('w')}><w:sdtContent>{paras}</w:sdtContent></w:sdt>"))
    d.add_paragraph("Body text")
    path = tmp_path / "toc.docx"
    d.save(str(path))
    return path


def test_entries_join_numbers_split_across_runs(tmp_path):
    path = _doc_with_toc(tmp_path, [("1.", "Ph.D. Program Requirements", ["3"]),
                                    ("3.", "Miscellaneous Information", ["1", "1"])])
    entries = toc_entries(docx.Document(str(path)))

    assert [t for t, _ in entries] == ["Ph.D. Program Requirements", "Miscellaneous Information"]
    assert ["".join(n.text for n in ns) for _, ns in entries] == ["3", "11"]


def test_locate_prefers_exact_heading_over_earlier_prose():
    pages = [
        ["Contents", "Final Public Oral (FPO): Department Instructions 8"],
        ["candidate is admitted to the final public oral (FPO) examination."],
        ["1.7 Final Public Oral (FPO): Department Instructions", "The readers..."],
    ]
    assert locate(["Final Public Oral (FPO): Department Instructions"], pages) == [3]


def test_locate_accepts_reworded_heading():
    pages = [["Contents"], ["Summary of first and second year requirements"]]
    assert locate(["First and Second Year requirements"], pages) == [2]


def test_locate_is_ordered_and_skips_contents_page():
    pages = [["Auditing Courses 10", "Leave of Absence 11"], ["Leave of Absence"], ["Auditing Courses"]]
    # "Auditing Courses" is listed first, so "Leave of Absence" is searched only after it.
    with pytest.raises(ValueError, match="Leave of Absence"):
        locate(["Auditing Courses", "Leave of Absence"], pages)
    assert locate(["Leave of Absence", "Auditing Courses"], pages) == [2, 3]


def test_locate_missing_heading_is_an_error():
    with pytest.raises(ValueError, match="Travel Support"):
        locate(["Travel Support"], [["Contents"], ["Nothing relevant here"]])


def test_update_rewrites_numbers(tmp_path, monkeypatch):
    path = _doc_with_toc(tmp_path, [("1.", "Ph.D. Program Requirements", ["3"]),
                                    ("3.", "Miscellaneous Information", ["1", "1"])])
    monkeypatch.setattr(toc, "pdf_page_lines", lambda _: [
        ["Contents"], ["Welcome"], ["1. Ph.D. Program Requirements"], ["..."], ["Miscellaneous Information"],
    ])

    assert update_toc_numbers(str(path), "unused.pdf") is True
    numbers = ["".join(n.text for n in ns) for _, ns in toc_entries(docx.Document(str(path)))]
    assert numbers == ["3", "5"]
    # Already correct: nothing to change, file untouched.
    assert update_toc_numbers(str(path), "unused.pdf") is False


def test_render_pdf_rerenders_until_stable(tmp_path, monkeypatch):
    from handbook_generator import converter as conv

    renders = []
    monkeypatch.setattr(conv.DocumentConverter, "convert", lambda self: renders.append(1))
    results = iter([True, False])
    monkeypatch.setattr(toc, "update_toc_numbers", lambda d, p: next(results))

    assert conv.render_pdf("a.docx", "a.pdf") is True
    assert len(renders) == 2


def test_render_pdf_fails_if_numbers_never_settle(monkeypatch):
    from handbook_generator import converter as conv

    monkeypatch.setattr(conv.DocumentConverter, "convert", lambda self: None)
    monkeypatch.setattr(toc, "update_toc_numbers", lambda d, p: True)
    with pytest.raises(RuntimeError, match="did not settle"):
        conv.render_pdf("a.docx", "a.pdf")
