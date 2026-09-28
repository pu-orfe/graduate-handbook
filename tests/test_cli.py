import os
import sys
import docx
import pytest
from bs4 import BeautifulSoup
from handbook_generator import cli
from handbook_generator.cli import main, get_copyright_year, EXIT_UNCHANGED

PAGE = """
<html><body>
  <div class="field--name-field-ps-body"><div class="tex2jax_process">
    <p>Welcome to the Department of ORFE.</p>
    <h2>Ph.D. Program Requirements</h2>
    <ul><li>Chair, Professor Test Chair, 609-000-0000</li></ul>
  </div></div>
  <footer>&copy; 2031 The Trustees of Princeton University</footer>
</body></html>
"""


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    """A working copy with a template and a fake scraper that never touches the network."""
    template = tmp_path / "template" / "graduate-handbook.docx"
    template.parent.mkdir()
    docx.Document().save(str(template))
    monkeypatch.chdir(tmp_path)

    page = {"html": PAGE}

    class FakeScraper:
        def __init__(self, url, media_dir):
            self.media_dir = media_dir

        def fetch_page(self):
            return BeautifulSoup(page["html"], "html.parser")

        def parse_contacts_and_download_images(self, soup):
            return {"chair_name": "Test Chair", "dgs_name": "Test DGS",
                    "chair_img_path": None, "dgs_img_path": None}

        def extract_main_content(self, soup):
            return soup.find("div", class_="field--name-field-ps-body")

    monkeypatch.setattr(cli, "HandbookScraper", FakeScraper)
    return {"root": tmp_path, "page": page}


def run(monkeypatch, *args):
    monkeypatch.setattr(sys, "argv", ["cli.py", "--docx-only", *args])
    with pytest.raises(SystemExit) as excinfo:
        main()
        raise SystemExit(0)
    return excinfo.value.code


def test_cli_help(mocker):
    mocker.patch.object(sys, "argv", ["cli.py", "--help"])
    with pytest.raises(SystemExit) as excinfo:
        main()
    assert excinfo.value.code == 0


def test_copyright_year():
    assert get_copyright_year(BeautifulSoup(PAGE, "html.parser")) == "2031"
    assert get_copyright_year(BeautifulSoup("<p>no footer</p>", "html.parser")) is None


def test_build_writes_outputs_and_snapshot(workspace, monkeypatch):
    root = workspace["root"]
    assert run(monkeypatch) == 0

    assert (root / "output" / "graduate-handbook.docx").exists()
    snapshot = (root / "snapshot" / "handbook.md").read_text()
    assert "year: 2031" in snapshot
    assert "## Ph.D. Program Requirements" in snapshot
    texts = [p.text for p in docx.Document(str(root / "output" / "graduate-handbook.docx")).paragraphs]
    assert any("Ph.D. Handbook 2031" in t for t in texts)


def test_skip_if_unchanged(workspace, monkeypatch):
    root = workspace["root"]
    assert run(monkeypatch, "--skip-if-unchanged") == 0
    docx_path = root / "output" / "graduate-handbook.docx"
    first_mtime = os.path.getmtime(docx_path)

    # Same page: reported as unchanged, and nothing is rebuilt.
    assert run(monkeypatch, "--skip-if-unchanged") == EXIT_UNCHANGED
    assert os.path.getmtime(docx_path) == first_mtime

    # Page edited: rebuilt and the snapshot follows.
    workspace["page"]["html"] = PAGE.replace("Welcome to", "A warm welcome to")
    assert run(monkeypatch, "--skip-if-unchanged") == 0
    assert "A warm welcome to" in (root / "snapshot" / "handbook.md").read_text()


def test_skip_if_unchanged_rebuilds_missing_outputs(workspace, monkeypatch):
    root = workspace["root"]
    assert run(monkeypatch) == 0
    (root / "output" / "graduate-handbook.docx").unlink()

    assert run(monkeypatch, "--skip-if-unchanged") == 0
    assert (root / "output" / "graduate-handbook.docx").exists()


def test_failed_build_leaves_snapshot_untouched(workspace, monkeypatch):
    root = workspace["root"]

    def boom(self, *a, **k):
        raise RuntimeError("builder exploded")

    monkeypatch.setattr(cli.HandbookBuilder, "build", boom)
    assert run(monkeypatch) == 1
    # Otherwise the next run would see "unchanged" and never retry.
    assert not (root / "snapshot" / "handbook.md").exists()


def test_fetch_failure_exits_nonzero(workspace, monkeypatch):
    def fail(self):
        raise ConnectionError("403 Forbidden")

    monkeypatch.setattr(cli.HandbookScraper, "fetch_page", fail)
    assert run(monkeypatch) == 1
    assert not (workspace["root"] / "snapshot" / "handbook.md").exists()


def test_missing_body_exits_nonzero(workspace, monkeypatch):
    workspace["page"]["html"] = "<html><body><p>Access denied</p></body></html>"
    assert run(monkeypatch) == 1
    assert not (workspace["root"] / "snapshot" / "handbook.md").exists()


def test_refuses_to_overwrite_template(workspace, monkeypatch):
    template = "template/graduate-handbook.docx"
    before = (workspace["root"] / template).read_bytes()
    assert run(monkeypatch, "--output-docx", template) == 1
    assert (workspace["root"] / template).read_bytes() == before


def test_contact_parse_failure_exits_nonzero(workspace, monkeypatch):
    def fail(self, soup):
        raise ValueError("Could not find the Chair on the page")

    monkeypatch.setattr(cli.HandbookScraper, "parse_contacts_and_download_images", fail)
    assert run(monkeypatch) == 1
    assert not (workspace["root"] / "snapshot" / "handbook.md").exists()


def test_page_disagreement_is_reported(workspace, monkeypatch, capsys):
    original = cli.HandbookScraper.parse_contacts_and_download_images

    def with_warning(self, soup):
        return dict(original(self, soup), warnings=["The page disagrees about the Chair"])

    monkeypatch.setattr(cli.HandbookScraper, "parse_contacts_and_download_images", with_warning)
    monkeypatch.setenv("GITHUB_ACTIONS", "true")
    assert run(monkeypatch) == 0
    assert "::warning::The page disagrees about the Chair" in capsys.readouterr().out


def test_findings_written_even_when_unchanged(workspace, monkeypatch):
    import json
    root = workspace["root"]
    assert run(monkeypatch) == 0
    # Unchanged page: the build is skipped, but checks still run (links can
    # break while the page stays the same).
    assert run(monkeypatch, "--skip-if-unchanged", "--findings", "reports/findings.json",
               "--skip-link-check") == EXIT_UNCHANGED
    result = json.loads((root / "reports" / "findings.json").read_text())
    assert "contacts" in result["checks_run"] and "links" not in result["checks_run"]
    assert isinstance(result["findings"], list)
