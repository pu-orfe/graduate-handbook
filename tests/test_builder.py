import os
import pytest
from bs4 import BeautifulSoup
import docx
from handbook_generator.builder import HandbookBuilder

@pytest.fixture
def sample_data():
    return {
        "chair_name": "Test Chair",
        "dgs_name": "Test DGS",
        "chair_img_path": None,
        "dgs_img_path": None
    }

@pytest.fixture
def sample_html():
    html = """
    <div class="tex2jax_process">
      <p>Welcome to the test handbook welcome letter!</p>
      <p>Go Tigers!</p>
      <h2>Ph.D. Program Requirements</h2>
      <p>Intro text.</p>
      <h3>Core courses</h3>
      <ul>
        <li>ORF 123 Course 1</li>
        <li>ORF 456 Course 2</li>
      </ul>
      <section class="cke-callout">
        <div class="cke-callout-content">
          <p>Must be approved 2 weeks prior.</p>
        </div>
      </section>
    </div>
    """
    return BeautifulSoup(html, "html.parser")

def test_builder_generates_docx(tmp_path, sample_data, sample_html):
    output_path = os.path.join(tmp_path, "output.docx")
    template_path = os.path.join(tmp_path, "template.docx")
    
    # Create a dummy template docx
    doc_temp = docx.Document()
    doc_temp.save(template_path)
    
    builder = HandbookBuilder(template_path=template_path, output_path=output_path, media_dir=str(tmp_path))
    builder.build(sample_data, sample_html, year="2026")
    
    assert os.path.exists(output_path)
    
    # Load generated document and inspect
    generated_doc = docx.Document(output_path)
    paragraphs = [p.text for p in generated_doc.paragraphs]
    
    # Check cover page text
    assert any("Ph.D. Handbook 2026" in p for p in paragraphs)
    
    # Check TOC title
    assert any("Contents" in p for p in paragraphs)
    
    # Check welcome heading
    assert any("Message from the Chair" in p for p in paragraphs)
    
    # Check body paragraph content
    assert any("welcome letter" in p for p in paragraphs)
    
    # Check heading mapping
    assert any("Ph.D. Program Requirements" in p for p in paragraphs)
    
    # Check callout message styling (contains IMPORTANT NOTICE)
    assert any("IMPORTANT NOTICE" in p for p in paragraphs)
